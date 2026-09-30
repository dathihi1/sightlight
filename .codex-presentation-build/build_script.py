from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'D:\duAnCaNhan\Sign_Light_Web')
OUT = ROOT / 'deliverables' / 'SignLight_EXE201_Outcome1_Presentation_Script_IV_V.docx'

NAVY = RGBColor(19,27,46)
TEAL = RGBColor(8,115,124)
MUTED = RGBColor(77,91,104)
PALE = 'EAF4F3'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72)
sec.bottom_margin = Inches(.68)
sec.left_margin = Inches(.78)
sec.right_margin = Inches(.78)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10.5)
normal.font.color.rgb = NAVY
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.16

title = styles['Title']
title.font.name = 'Arial'
title.font.size = Pt(25)
title.font.bold = True
title.font.color.rgb = RGBColor(0,0,0)
title.paragraph_format.space_after = Pt(9)

for nm,sz,after in [('Heading 1',15,7),('Heading 2',11.5,4)]:
    st=styles[nm]
    st.font.name='Arial'; st.font.size=Pt(sz); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(17 if nm=='Heading 1' else 9)
    st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True

cue_style=styles.add_style('Stage Cue', WD_STYLE_TYPE.PARAGRAPH)
cue_style.base_style=normal
cue_style.font.name='Arial'; cue_style.font.size=Pt(9.5); cue_style.font.italic=True; cue_style.font.color.rgb=TEAL
cue_style.paragraph_format.space_after=Pt(5)
cue_style.paragraph_format.keep_with_next=True

meta_style=styles.add_style('Slide Meta', WD_STYLE_TYPE.PARAGRAPH)
meta_style.base_style=normal
meta_style.font.name='Arial'; meta_style.font.size=Pt(9); meta_style.font.bold=True; meta_style.font.color.rgb=MUTED
meta_style.paragraph_format.space_after=Pt(5)
meta_style.paragraph_format.keep_with_next=True

def cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def cell_border(cell):
    tcPr=cell._tc.get_or_add_tcPr()
    b=OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); b.append(el)
    tcPr.append(b)
def table(rows, widths):
    t=doc.add_table(rows=1, cols=len(widths)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,h in enumerate(rows[0]):
        c=t.rows[0].cells[i]; c.text=h; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell_shading(c,'203148');cell_border(c)
        for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9.5)
    for j,row in enumerate(rows[1:]):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i];c.text=val;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell_border(c)
            if j%2:cell_shading(c,PALE)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs:r.font.size=Pt(9.4)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def para(text='', style=None):
    p=doc.add_paragraph(style=style)
    p.add_run(text)
    return p
def slide(n,title,speaker,time,cue,speech):
    doc.add_heading(f'Slide {n}  {title}',level=1)
    para(f'Speaker: {speaker}    Target: {time}', 'Slide Meta')
    if cue: para(f'Stage cue: {cue}', 'Stage Cue')
    for part in speech:
        para(part)

doc.add_paragraph('SignLight EXE201 Outcome 1 Presentation Script', 'Title')
para('English speaking script for the 17-slide presentation, including Parts IV and V. Only Lê Tùng Dương, Nguyễn Thế Duy and Phan Bùi Bá Đạt speak. Phạm Việt Hoàng and Hoàng Xuân Thọ remain listed on the team slide for their finance, business and marketing responsibilities.')
para('Run time: approximately 23 minutes for slides and live actions, followed by about 2 minutes of questions. The slide deck and this script should be rehearsed together; demo pauses are included in the target times.')

doc.add_heading('Speaker and timing plan',level=1)
table([
 ['Speaker','Slides','Approximate time'],
 ['Nguyễn Thế Duy','1, 2, 5, 6, 8, 12, 13, 14, 15, 16, 17','11 min 40 sec'],
 ['Lê Tùng Dương','3, 4, 9, 11','5 min 55 sec'],
 ['Phan Bùi Bá Đạt','7, 10','5 min 25 sec'],
],[1.66,2.7,2.35])

doc.add_heading('Before presenting',level=1)
for x in [
 'Create a dedicated demo account, sign in once, and keep credentials private. Do not show a password on screen.',
 'Open the published site at https://sightlight1.vercel.app/ and test /hoc, /luyen-ai, /tu-dien and /nang-cap on the presentation laptop.',
 'Check webcam access, lighting, the chosen target sign, its teaching video and the API response. Have a second browser tab ready with the product home page.',
 'Do not make a real payment during the demonstration. The pricing page is sufficient for Outcome 1.',
 'Google Analytics 4 is a planned addition. Do not say it is already installed or report traffic figures that have not been measured.',
 'The Part IV Early Bird lifetime price, B2B discounts and uptime SLA are proposals requiring finance, legal and technical review; they are not live product terms.'
]:
    p=para('• '+x);p.paragraph_format.left_indent=Inches(.18)

doc.add_heading('Short answers for likely questions',level=1)
table([
 ['Question','Suggested answer'],
 ['Is the AI accurate?','We have a working recognition flow and local ONNX model files. We have not completed a real webcam accuracy study, so we do not claim a final accuracy figure today.'],
 ['Is the website public?','Yes. The published MVP is available at sightlight1.vercel.app. Camera practice requires a signed-in account.'],
 ['Are the prices final?','No. They are proposed EXE201 prices represented in the current product code and will be reviewed against costs and user response.'],
 ['Is analytics installed?','Not yet confirmed in the current frontend. GA4 instrumentation is in the W3 plan.'],
 ['Are the sales and media results achieved?','The figures on the new slides are targets from the plan. We will report outcomes only with survey records, backend registrations, GA4 data and signed documents.']
],[1.75,4.95])

slide(1,'SignLight','Nguyễn Thế Duy','0:40','Open on the title slide. Speak slowly and make eye contact.',[
 'Good morning. We are the SignLight team. SignLight is a web product for learning Vietnamese Sign Language. Our core idea is simple: a learner can watch a model sign, practice in front of a camera, and receive feedback about the sign they performed.',
 'Today we will introduce our team, show the product that is already online, demonstrate the main learning flow, and explain what we will complete during the remaining seven weeks.'
])

slide(2,'Today’s presentation','Nguyễn Thế Duy','0:35','Point briefly to the live demo block.',[
 'We will keep this presentation within the 20 to 25 minute limit. The short project overview and team plan come first. Then we will spend the largest block on the product itself, especially the Camera AI function. We will finish with the development schedule, sales kit and media plan, then leave time for questions.'
])

slide(3,'The learning problem','Lê Tùng Dương','1:15','Take over from Duy. Pause after the first sentence.',[
 'A video can show a sign very clearly. But when someone practices alone, the video cannot tell them whether their own movement was correct. That gap matters because hand position, movement and timing are all part of the sign.',
 'Our learning flow responds to that gap. First, the learner watches a VSL model video. Next, they perform the sign in front of the camera. Finally, SignLight shows a recognized label and guidance. The purpose is to give learners a useful next step when they need to try again.',
 'My role is to develop lesson content, quizzes and teaching videos. The camera feedback becomes much more useful when the target sign and its lesson are carefully prepared.'
])

slide(4,'Published MVP','Lê Tùng Dương','1:05','Open the website in a separate tab only if the network is stable; otherwise keep the link visible on the slide.',[
 'The current web application is published at sightlight1.vercel.app. Visitors can reach the landing page, and the product contains routes for learning, dictionary search, camera practice and Premium plans.',
 'For Outcome 1, our main demonstration is the learner journey rather than a tour of every page. We will start with a lesson, move to the camera, then look up a sign in the dictionary. That sequence shows how teaching content and the AI feature work together.',
 'The product is still an MVP. Published pages do not mean that every planned lesson or every model quality target has been completed.'
])

slide(5,'Team roles and deliverables','Nguyễn Thế Duy','1:00','Keep this factual; Hoàng and Thọ do not take a speaking turn.',[
 'We have five members with distinct work areas. Dương leads the learning content, quizzes and teaching videos. Hoàng and Thọ handle finance, business planning and marketing. I develop the landing page, backend and product interface. Đạt trains the AI model and builds the camera AI function.',
 'Three of us are speaking today: Dương for the learning experience, Đạt for the AI demonstration, and me for the product, proposed plans and delivery timeline. The team slide still reflects all five contributors.'
])

slide(6,'Product catalog and proposed prices','Nguyễn Thế Duy','1:30','Point to Free, then move from left to right across the Premium options.',[
 'Our proposed catalog has a Free entry level and three Premium durations. Free lets a learner start with the first learning unit and a limited number of AI practice attempts each day. The monthly Premium plan is ninety-nine thousand đồng. Six months is four hundred and ninety-nine thousand đồng, and one year is eight hundred and ninety-nine thousand đồng.',
 'The Premium descriptions in the current product code cover the full course and unlimited AI practice. These are proposed EXE201 prices, not validated sales results. Hoàng and Thọ will test the business assumptions, costs and marketing response before we treat the price structure as final.',
 'On the live site we will show the pricing page, but we will not submit a payment during this presentation.'
])

slide(7,'How camera practice works','Phan Bùi Bá Đạt','1:25','Use your hands to trace the four steps rather than reading every line.',[
 'Here is the camera flow we have implemented. The learner watches a target sign video and records one complete movement. The browser reads the webcam frames and extracts numeric hand and body features. The recognition model can run locally in the browser through ONNX WebAssembly. The numeric features are also sent to the backend so the attempt and daily quota can be tracked.',
 'A key privacy boundary is that camera pixels stay on the learner’s device. The request contains a tensor of numbers rather than an image or video file. The current code also blocks image fields in the inference payload.',
 'We will demonstrate the interaction next. We are still measuring real webcam accuracy and checking the server model deployment, so today we will not claim a final recognition accuracy number.'
])

slide(8,'Live demo journey','Nguyễn Thế Duy','0:25','Switch to the signed-in browser tab and hand control to Dương.',[
 'Now we will use the published product. We will open a lesson, practice one sign with the camera, and finish with the dictionary and pricing page. Dương will start with the learning path.'
])

slide(9,'Demo lesson and learning path','Lê Tùng Dương','2:10','Navigate to /hoc. Open an available lesson, play a working video, answer one question, and show progress if available.',[
 'This is the learner’s course path. We choose an available lesson and open the sign video. The video gives the learner a visual model before they try to reproduce the movement.',
 'Here is one lesson question. The response is checked by the backend, and lesson progress can be saved after completion. In this demo, the important point is the sequence: see the sign, understand the task, and practice it. We are continuing to prepare lesson scripts, quizzes and teaching videos so this learning path has strong instructional content, not only working screens.',
 'I will now hand over to Đạt for the main Camera AI demonstration.'
])

slide(10,'Demo Camera AI','Phan Bùi Bá Đạt','4:00','Navigate to /luyen-ai. Allow camera access only on the presenter’s device. Choose a target sign, play its model video, record a complete movement, then wait for the response.',[
 'This is our main product demonstration. The target sign appears next to the camera view. I will watch the model video, position myself in the camera frame and perform one complete movement. The browser extracts the motion features on this device. It does not upload the camera video.',
 'I am starting the recording now. After I finish, SignLight shows the sign it recognized, a confidence value and alternative labels when available. If the target does not match, the learner can use the feedback and try again. The free tier also has a daily practice limit, which the backend tracks.',
 'Please read this result as a live product response, not as proof of overall model accuracy. Accuracy requires a separate test with multiple people and lighting conditions. We are doing that work in the next phase. The current server-side model deployment also needs validation, so any simulated server response must be identified as such.'
])

doc.add_heading('Camera demo fallback lines',level=2)
para('If camera permission is denied: “The browser did not grant camera access. This is a device permission issue, so I will show the practice screen and explain the recording flow without presenting a recognition result.”')
para('If the API is slow or unavailable: “The live service did not return a response in time. The learning interface and local camera flow are visible; we will not invent an AI result. We will continue with the dictionary and our validation plan.”')
para('If a stub warning appears: “This server response is marked as simulation. It demonstrates the interaction, but it is not an accuracy measurement.”')

slide(11,'Demo dictionary and upgrade path','Lê Tùng Dương','1:25','Navigate to /tu-dien; search one prepared term. Then briefly show /nang-cap.',[
 'After practice, the learner can search the VSL dictionary. We can enter a Vietnamese term with or without diacritics and open its sign video. This makes the dictionary useful both before a lesson and when a learner wants to review a sign afterward.',
 'The upgrade page shows the proposed Premium options that Duy introduced earlier. We will stop on the plan page. This presentation does not require a payment transaction. The important point is that the learning and practice flow has a clear entry point and a proposed way to support continued development.'
])

slide(12,'MVP readiness','Nguyễn Thế Duy','1:20','Return from the browser to the deck. Do not rush the limitations.',[
 'The site is online, and we can show the landing page, learning route, dictionary route, camera practice interface and Premium plan interface. That is the evidence we have today.',
 'Before the Outcome 1 presentation, we still need a dedicated demo account, a full camera and API rehearsal, and a real webcam evaluation. The repository also does not yet show Google Analytics 4 instrumentation. We will add it so the team can measure visits, sign-ups and important learning events. Hoàng and Thọ can then use actual website performance data in later business and marketing reports.',
 'We are separating a published MVP from a fully validated release. That keeps our demonstration credible and gives us a clear work plan.'
])

slide(13,'Seven-week development plan','Nguyễn Thế Duy','1:30','Point to each phase in sequence: W3, W6, W8.',[
 'We have divided the remaining seven weeks into three phases, matching the guideline milestones. By week three, Version 1.0 is a stable MVP: a tested demo account, reliable camera and API flow, a pass on the lesson content, and Google Analytics page and event tracking.',
 'By week six, Version 2.0 adds more learning depth. The focus is quizzes, better AI feedback, checkout quality assurance and a real webcam test across different people and lighting conditions.',
 'The final phase runs from week six to week eight. We aim to launch the fuller product publicly in week six or seven, then use the remaining time to improve content, retention and issues found in testing. The website is already online; this milestone refers to a validated product release rather than the first publication of a landing page.'
])

slide(14,'Function list by version','Nguyễn Thế Duy','1:20','Summarize priorities; do not read every function verbatim.',[
 'This function list turns the timeline into concrete scope. Version 1.0 focuses on what a learner must be able to do during the demo: sign in, follow a course, open lessons, use the dictionary and practice with the camera. We add analytics instrumentation so later decisions are based on usage data.',
 'Version 2.0 strengthens the learning loop with quizzes, improved guidance, payment testing and webcam evaluation. Version 3.0 is about release quality: content review, accessibility, privacy, stability and a final launch checklist.',
 'This order reflects urgency. We first make the existing journey dependable, then improve depth, and finally polish and validate the wider release.'
])

slide(15,'Sales plan and B2B kit','Nguyễn Thế Duy','1:20','Point to the kit on the left, then the two sales actions on the right. Hoàng and Thọ do not speak.',[
 'Part IV turns our business idea into a sales package. The proposed B2B kit contains a company brochure and LMS catalog, a tiered quote for education buyers, an ESG and CSR partnership proposal with measurable impact, and a service agreement with a support and uptime target. These materials give a school, center or sponsor a clear basis for discussion.',
 'The document proposes a thirty to fifty percent education discount and a ninety-nine point nine percent uptime commitment. Both require cost, legal and infrastructure review before we offer them. It also proposes an Early Bird lifetime package at one hundred and ninety-nine thousand đồng for the first one hundred buyers. This is a campaign concept, not a live price on the current website.',
 'For B2B outreach in weeks four to six, Hoàng and Thọ plan to contact fifteen inclusive education centers and ten food, beverage or retail chains. We will track sent kits, meetings, proposals and signed agreements separately.'
])

slide(16,'Media master plan','Nguyễn Thế Duy','1:30','Move through the four periods in order. State that the numbers are targets.',[
 'Part V maps communication from week one to week eight. In weeks one and two, the plan starts with a Facebook fanpage and beta user research, targeting at least sixty-five valid surveys and eighteen interviews. In weeks three and four, short videos and the free VSL dictionary aim for more than two thousand website visits and three hundred registrations.',
 'Weeks five and six add the Light Ambassador campus campaign and communication around Version 2.0 and the B2B LMS offer. The targets are one thousand registered accounts and one or two B2B memoranda of understanding. Weeks seven and eight focus on an impact report and partner case studies, aiming for five thousand visitors and the revenue target.',
 'These are targets from the plan. We will verify actual results using survey records, signed documents, backend account counts and analytics once GA4 is installed. Hoàng and Thọ own the marketing work; I am presenting their plan today.'
])

slide(17,'Closing','Nguyễn Thế Duy','0:30','Pause on the live URL, then invite questions.',[
 'SignLight is already available as a web MVP. Our main contribution is the path from watching a VSL sign to practicing it and seeing feedback. The next seven weeks will turn that working journey into a better tested and more complete learning product. Thank you. We welcome your questions.'
])

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('SignLight  |  EXE201 Outcome 1')
run.font.name='Arial';run.font.size=Pt(8);run.font.color.rgb=MUTED

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
